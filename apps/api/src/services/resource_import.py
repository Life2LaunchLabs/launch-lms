"""Bring a spreadsheet of links in as resources: read the file, match its columns, check the rows, then import."""

import csv
import io
import re
from urllib.parse import urlparse
from uuid import uuid4

from fastapi import HTTPException, Request, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy import func
from sqlmodel import Session, select
from src.db.resources import Resource, ResourceChannel, ResourceChannelResource, ResourceTag, ResourceTagLink, ResourceTypeEnum
from src.db.users import PublicUser
from src.services import resources as base

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
FIELDS = ("link", "title", "description", "type", "provider", "minutes", "channel", "tags")

# Header words that suggest a field, most specific first; a header matches the first field any word appears in.
HEADER_HINTS: list[tuple[str, tuple[str, ...]]] = [
    ("link", ("link", "url", "href", "website", "web address")),
    ("minutes", ("minutes", "mins", "estimated time", "est time", "duration", "time", "length")),
    ("description", ("description", "summary", "about", "details", "overview")),
    ("type", ("type", "format", "kind", "media")),
    ("provider", ("provider", "source", "publisher", "author", "site")),
    ("channel", ("channel", "category", "pathway", "collection", "folder", "group", "section")),
    ("tags", ("tags", "tag", "keywords", "topics", "labels")),
    ("title", ("title", "name", "resource", "heading")),
]
TYPE_WORDS = {item.value: item.value for item in ResourceTypeEnum} | {
    "articles": "article", "videos": "video", "tools": "tool", "guides": "guide", "courses": "course",
    "assessments": "assessment", "quiz": "assessment", "test": "assessment", "website": "article", "page": "article",
    "youtube": "video", "app": "tool", "lesson": "course", "worksheet": "guide",
}
LIST_SPLIT = re.compile(r"\s*[|;]\s*")


def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _read_rows(filename: str, raw: bytes) -> list[list[str]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
        except Exception:
            raise HTTPException(status_code=400, detail="That spreadsheet couldn't be opened. Save it as .xlsx or .csv and try again.") from None
        return [[_clean(cell) for cell in row] for row in sheet.iter_rows(values_only=True, max_row=MAX_ROWS + 1)]
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise HTTPException(status_code=400, detail="That file isn't text we can read. Save it as .csv or .xlsx.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [[_clean(cell) for cell in row] for row, _ in zip(csv.reader(io.StringIO(text), dialect), range(MAX_ROWS + 1))]


def guess_field(header: str, taken: set[str]) -> str:
    """The field a header most likely means, or "skip"; each field is guessed for one column at most."""
    words = re.sub(r"[^a-z0-9]+", " ", header.lower()).strip()
    for field, hints in HEADER_HINTS:
        if field not in taken and any(re.search(rf"\b{re.escape(hint)}\b", words) for hint in hints):
            return field
    return "skip"


async def read_spreadsheet(request: Request, org_id: int, file: UploadFile, current_user: PublicUser, db_session: Session) -> dict:
    """The file's columns with a guessed field and an example each, and its rows, for matching on the page."""
    base._get_org_or_404(org_id, db_session)
    base._require_staff(current_user, org_id, db_session)
    name = file.filename or "spreadsheet.csv"
    if not name.lower().endswith((".csv", ".tsv", ".txt", ".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Choose a .csv or .xlsx file.")
    raw = await file.read(MAX_FILE_BYTES + 1)
    if len(raw) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="That file is over 5 MB. Split it into smaller files.")
    rows = [row for row in _read_rows(name, raw) if any(row)]
    if len(rows) < 2:
        raise HTTPException(status_code=400, detail="The file needs a header row and at least one row of links.")
    header, body = rows[0], rows[1 : MAX_ROWS + 1]
    width = max(len(row) for row in rows)
    taken: set[str] = set()
    columns = []
    for index in range(width):
        label = header[index] if index < len(header) and header[index] else f"Column {index + 1}"
        field = guess_field(label, taken)
        if field == "skip" and "link" not in taken and any(_link_of(row[index] if index < len(row) else "") for row in body[:5]):
            field = "link"
        if field != "skip":
            taken.add(field)
        example = next((row[index] for row in body if index < len(row) and row[index]), "")
        columns.append({"name": label, "field": field, "example": example[:140]})
    return {
        "file_name": name,
        "columns": columns,
        "rows": [row + [""] * (width - len(row)) for row in body],
        "truncated": len(rows) - 1 > MAX_ROWS,
    }


class ImportRequest(BaseModel):
    """Rows already matched to fields on the page ({"link": ..., "title": ...}), and where they go."""

    rows: list[dict[str, str]] = Field(max_length=MAX_ROWS)
    destination: str = "column"  # "column": a channel per Channel value; "one": all in channel_uuid or a new channel_name; "none"
    channel_uuid: str | None = None
    channel_name: str | None = None
    dry_run: bool = True


def _link_of(value: str) -> str | None:
    value = value.strip()
    if not value or " " in value:
        return None
    if not re.match(r"^https?://", value, re.IGNORECASE):
        if not re.match(r"^(www\.)?[a-z0-9-]+(\.[a-z0-9-]+)+(/|$)", value, re.IGNORECASE):
            return None
        value = f"https://{value}"
    parsed = urlparse(value)
    return value if parsed.scheme in ("http", "https") and "." in parsed.netloc else None


def _minutes(value: str) -> int | None:
    match = re.search(r"\d+(?:\.\d+)?", value or "")
    if not match:
        return None
    number = float(match.group())
    if re.search(r"\b(h|hr|hrs|hour|hours)\b", value.lower()):
        number *= 60
    return int(round(number)) or None


def _type_of(value: str) -> str:
    word = (value or "").strip().lower()
    return TYPE_WORDS.get(word) or next((TYPE_WORDS[key] for key in TYPE_WORDS if key in word.split()), "other")


def _plan_rows(org_id: int, rows: list[dict[str, str]], db_session: Session) -> list[dict]:
    existing = {url.rstrip("/") for url in db_session.exec(select(Resource.external_url).where(Resource.org_id == org_id)).all()}
    seen: set[str] = set()
    planned = []
    for index, row in enumerate(rows):
        link = _link_of(row.get("link", ""))
        title = row.get("title", "").strip()
        item = {
            "row": index + 1,
            "link": link or row.get("link", "").strip(),
            "title": title[:300] or (urlparse(link).netloc.removeprefix("www.") if link else ""),
            "type": _type_of(row.get("type", "")),
            "channels": [name[:120] for name in LIST_SPLIT.split(row.get("channel", "").strip()) if name],
            "status": "ready",
        }
        if not link:
            item["status"], item["error"] = "invalid", "No link" if not item["link"] else "Not a web link"
        elif link.rstrip("/") in existing:
            item["status"] = "duplicate"
        elif link.rstrip("/") in seen:
            item["status"] = "repeat"
        else:
            seen.add(link.rstrip("/"))
        item["_row"] = row
        item["_titled"] = bool(title)
        planned.append(item)
    return planned


def _org_channel(org_id: int, name: str, db_session: Session) -> ResourceChannel | None:
    return db_session.exec(select(ResourceChannel).where(ResourceChannel.org_id == org_id, func.lower(ResourceChannel.name) == name.lower())).first()


def _org_tags(org_id: int, names: set[str], db_session: Session) -> dict[str, ResourceTag]:
    tags = {tag.name.lower(): tag for tag in db_session.exec(select(ResourceTag).where(ResourceTag.org_id == org_id)).all()}
    for name in names:
        if name.lower() not in tags:
            tag = ResourceTag(org_id=org_id, name=name, tag_uuid=f"tag_{uuid4()}", creation_date=base._now(), update_date=base._now())
            db_session.add(tag)
            tags[name.lower()] = tag
    db_session.flush()
    return tags


async def import_rows(request: Request, org_id: int, data: ImportRequest, current_user: PublicUser, db_session: Session) -> dict:
    """Check the rows (dry_run) or create the ready ones and put them in their channels. Links already in the org are skipped."""
    base._get_org_or_404(org_id, db_session)
    base._require_staff(current_user, org_id, db_session)
    if data.destination not in ("column", "one", "none"):
        raise HTTPException(status_code=422, detail="Unknown destination")
    target = None
    if data.destination == "one" and data.channel_uuid:
        target = base._get_channel_or_404(data.channel_uuid, db_session)
        if target.org_id != org_id:
            raise HTTPException(status_code=400, detail="That channel isn't your organization's")
    elif data.destination == "one" and not (data.channel_name or "").strip():
        raise HTTPException(status_code=422, detail="Name the channel they go in")
    planned = _plan_rows(org_id, data.rows, db_session)
    ready = [item for item in planned if item["status"] == "ready"]
    channel_names: dict[str, str] = {}  # each name as first written, by its lowercase
    for name in (name for item in ready for name in item["channels"]) if data.destination == "column" else ():
        channel_names.setdefault(name.lower(), name)
    channel_names = list(channel_names.values())
    new_channels = sorted(name for name in channel_names if not _org_channel(org_id, name, db_session))
    if data.destination == "one" and not target:
        new_channels = [] if _org_channel(org_id, data.channel_name.strip(), db_session) else [data.channel_name.strip()]
    summary = {
        "rows": [{key: value for key, value in item.items() if not key.startswith("_")} for item in planned],
        "ready": len(ready),
        "duplicates": sum(item["status"] in ("duplicate", "repeat") for item in planned),
        "invalid": sum(item["status"] == "invalid" for item in planned),
        "new_channels": new_channels,
        "channel_count": len(channel_names) if data.destination == "column" else int(data.destination == "one"),
    }
    if data.dry_run:
        return {**summary, "created": 0}

    channels: dict[str, ResourceChannel] = {}

    def channel_named(name: str) -> ResourceChannel:
        key = name.lower()
        if key not in channels:
            channel = _org_channel(org_id, name, db_session)
            if not channel:
                channel = ResourceChannel(org_id=org_id, channel_uuid=f"channel_{uuid4()}", name=name, creation_date=base._now(), update_date=base._now())
                db_session.add(channel)
                db_session.flush()
            channels[key] = channel
        return channels[key]

    next_order: dict[int, int] = {}

    def link(channel: ResourceChannel, resource: Resource) -> None:
        if channel.id not in next_order:
            last = db_session.exec(select(func.max(ResourceChannelResource.sort_order)).where(ResourceChannelResource.channel_id == channel.id)).one()
            next_order[channel.id] = (last or 0) + 1
        db_session.add(ResourceChannelResource(channel_id=channel.id, resource_id=resource.id, sort_order=next_order[channel.id], creation_date=base._now(), update_date=base._now()))
        next_order[channel.id] += 1

    tags = _org_tags(org_id, {name[:60] for item in ready for name in LIST_SPLIT.split(item["_row"].get("tags", "").strip()) if name}, db_session)
    one = target or (channel_named(data.channel_name.strip()) if data.destination == "one" else None)
    created: list[tuple[int, bool, bool]] = []
    for item in ready:
        row = item["_row"]
        description = row.get("description", "").strip() or None
        resource = Resource(
            org_id=org_id,
            created_by_user_id=current_user.id,
            resource_uuid=f"resource_{uuid4()}",
            title=item["title"],
            description=description,
            resource_type=item["type"],
            provider_name=row.get("provider", "").strip()[:200] or urlparse(item["link"]).netloc.removeprefix("www.") or None,
            external_url=item["link"],
            estimated_time=_minutes(row.get("minutes", "")),
            creation_date=base._now(),
            update_date=base._now(),
        )
        db_session.add(resource)
        db_session.flush()
        for name in {name[:60].lower() for name in LIST_SPLIT.split(row.get("tags", "").strip()) if name}:
            db_session.add(ResourceTagLink(resource_id=resource.id, tag_id=tags[name].id))
        for channel in [one] if one else [channel_named(name) for name in dict.fromkeys(item["channels"])]:
            link(channel, resource)
        created.append((resource.id, not item["_titled"], description is None))
    db_session.commit()
    return {**summary, "created": len(created), "_fill_in": created}


def fill_in_imported(org_id: int, created: list[tuple[int, bool, bool]]) -> None:
    """After an import: each new resource's picture, and its title and description when the sheet had none."""
    from src.services.demo.namespaces import context_session

    with context_session() as db_session:
        for resource_id, fill_title, fill_description in created:
            resource = db_session.get(Resource, resource_id)
            if not resource:
                continue
            found = base.enrich_resource_metadata(resource.external_url)
            resource.cover_image_url = resource.cover_image_url or found.get("cover_image_url")
            if fill_title and found.get("title"):
                resource.title = found["title"][:300]
            if fill_description and found.get("description"):
                resource.description = found["description"]
            db_session.add(resource)
            db_session.commit()
